#!/usr/bin/env python3
"""Generate the colleague-facing six-column feedback workbook."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


HEADERS = ["编号", "能力", "实际问题", "期望结果", "严重度", "证据"]
EXAMPLE = [
    "P02-004",
    "Table Shell",
    "主要终点表没有写分析人群和时点",
    "表头或脚注明确 ITT、Week 12",
    "High",
    "P02-SHELL-R1",
]
INSTRUCTIONS = (
    "填写说明：第 2 行为示例，请从第 4 行开始填写。每行记录一个问题；"
    "能力和严重度请用下拉框选择；证据可留空；请勿填写姓名、患者信息或敏感资料。"
)
WIDTHS = [14, 18, 48, 48, 12, 28]


def style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[3]:
        cell.fill = fill
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A3:F500"
    ws.row_dimensions[3].height = 30


def build(output: Path) -> None:
    wb = Workbook()
    feedback = wb.active
    feedback.title = "反馈"

    feedback.merge_cells("A1:F1")
    feedback["A1"] = INSTRUCTIONS
    feedback["A1"].fill = PatternFill("solid", fgColor="EAF2FF")
    feedback["A1"].font = Font(name="Arial", color="1E3A5F")
    feedback["A1"].alignment = Alignment(vertical="center", wrap_text=True)
    feedback.row_dimensions[1].height = 42

    for index, value in enumerate(EXAMPLE, start=1):
        cell = feedback.cell(row=2, column=index, value=value)
        cell.fill = PatternFill("solid", fgColor="F3F4F6")
        cell.font = Font(name="Arial", color="4B5563", italic=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    feedback.row_dimensions[2].height = 34

    for index, value in enumerate(HEADERS, start=1):
        feedback.cell(row=3, column=index, value=value)

    for index, width in enumerate(WIDTHS, start=1):
        feedback.column_dimensions[chr(64 + index)].width = width
    style_header(feedback)

    capability = DataValidation(
        type="list",
        formula1='"CRF Review,SAP,Table Shell,Workflow,Other"',
        allow_blank=False,
    )
    severity = DataValidation(
        type="list",
        formula1='"Blocker,High,Normal,Low"',
        allow_blank=False,
    )
    feedback.add_data_validation(capability)
    feedback.add_data_validation(severity)
    capability.add("B4:B500")
    severity.add("E4:E500")
    for row in feedback.iter_rows(min_row=4, max_row=500, min_col=1, max_col=6):
        for cell in row:
            cell.font = Font(name="Arial", color="1F2937")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    feedback.sheet_view.showGridLines = False

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: make-feedback-template.py <output.xlsx>")
    build(Path(sys.argv[1]))
